import os
import sys
import sqlite3
from datetime import datetime
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

DB_PATH = os.path.join(os.path.dirname(__file__), "warehouse.db")
EXCEL_PATH = r"D:\Downloads\KingsMan (1).xlsx"

def clean_code(val):
    if not val:
        return ""
    s = str(val).strip()
    if s.startswith(('=', 'Tìm', 'Tổng', 'None')):
        return ""
    return s

def guess_category_and_info(code):
    c_upper = code.upper()
    cat = "Giày Da"
    color = "Đen"
    material = "Da bò cao cấp"

    if c_upper.startswith("CT"):
        cat = "Oxford Cap-toe"
        if "D" in c_upper: color = "Đen"
        elif "N" in c_upper: color = "Nâu"
        elif "S" in c_upper: color = "Nâu Sáp"; material = "Da bò sáp"
        elif "T" in c_upper: color = "Nâu Trầm"
        elif "V" in c_upper: color = "Vàng Bò"
    elif c_upper.startswith("BR"):
        cat = "Derby Brogue"
        if "GỖ" in c_upper: material = "Da bò đế phíp gỗ"
        if "D" in c_upper: color = "Đen"
        elif "N" in c_upper: color = "Nâu"
        elif "V" in c_upper: color = "Vàng Bò"
        elif "L" in c_upper: color = "Nâu Lợt"
    elif c_upper.startswith("L"):
        cat = "Penny Loafer"
        if "HB" in c_upper: material = "Da bò hạt"
        if "D" in c_upper: color = "Đen"
        elif "N" in c_upper: color = "Nâu"
        elif "V" in c_upper: color = "Vàng Bò"
    elif c_upper.startswith("CHEL"):
        cat = "Chelsea Boot"
        if "N" in c_upper: color = "Nâu"; material = "Da bò sáp/Nappa"
        else: color = "Đen"
    elif c_upper.startswith("DT"):
        cat = "Derby Trơn"
        if "N" in c_upper: color = "Nâu"
        else: color = "Đen"
    elif c_upper.startswith("DB"):
        cat = "Derby Bò"
        if "N" in c_upper: color = "Nâu"
        elif "M" in c_upper: color = "Mộc"
    elif c_upper.startswith("MTS"):
        cat = "Monkstrap"
        if "N" in c_upper: color = "Nâu"
        else: color = "Đen"
    elif "DÂY LƯNG" in c_upper:
        cat = "Phụ Kiện Da"
        material = "Da bò nguyên tấm"
        color = "Đen/Nâu"

    name = f"Giày Da Kingsman {code}" if cat != "Phụ Kiện Da" else f"Thắt Lưng Da Kingsman {code}"
    return cat, name, color, material

def run_import():
    print("==================================================================")
    print(f"BẮT ĐẦU NHẬP DỮ LIỆU TỪ: {EXCEL_PATH}")
    print("==================================================================")

    if not os.path.exists(EXCEL_PATH):
        print(f"LỖI: Không tìm thấy file {EXCEL_PATH}!")
        return

    wb_val = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    
    # Kết nối SQLite
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute("PRAGMA foreign_keys = OFF;")

    # Xóa sạch bảng để nạp dữ liệu chuẩn mới
    cursor.execute("DELETE FROM transaction_items;")
    cursor.execute("DELETE FROM transactions;")
    cursor.execute("DELETE FROM inventory;")
    cursor.execute("DELETE FROM product_variants;")
    cursor.execute("DELETE FROM products;")
    cursor.execute("DELETE FROM warehouses;")
    
    # 1. Khởi tạo 2 Kho chuẩn theo mô tả người dùng:
    # Kho 1 = Kho Tổng (Ở Nhà)
    # Kho 2 = Cửa Hàng (Showroom)
    cursor.execute("""
    INSERT INTO warehouses (id, code, name, address, phone, is_default) VALUES
    (1, 'KHO-TONG', 'Kho 1 - Kho Tổng (Ở Nhà)', 'Kho Tổng / Lưu Trữ Tại Nhà', '0988.xxx.xxx', 1),
    (2, 'KHO-STORE', 'Kho 2 - Cửa Hàng (Showroom)', 'Cửa Hàng Trưng Bày & Bán Lẻ', '0977.xxx.xxx', 0);
    """)

    # 2. Lấy danh sách giá từ sheet Nhập và Xuất
    prices = {} # code -> {cost, retail}
    if 'Nhập' in wb_val.sheetnames:
        sh_n = wb_val['Nhập']
        for r in range(2, sh_n.max_row + 1):
            code = clean_code(sh_n.cell(row=r, column=2).value)
            price = sh_n.cell(row=r, column=5).value
            if code and isinstance(price, (int, float)) and price > 0:
                p_val = price * 1000 if price < 10000 else price
                if code not in prices: prices[code] = {}
                prices[code]['cost'] = p_val

    if 'Xuất' in wb_val.sheetnames:
        sh_x = wb_val['Xuất']
        for r in range(2, sh_x.max_row + 1):
            code = clean_code(sh_x.cell(row=r, column=2).value)
            price = sh_x.cell(row=r, column=5).value
            if code and isinstance(price, (int, float)) and price > 0:
                p_val = price * 1000 if price < 10000 else price
                if code not in prices: prices[code] = {}
                prices[code]['retail'] = p_val

    # 3. Thu thập dữ liệu tồn kho từ sheet 'Kho' (Cửa Hàng) và sheet 'Ở Nhà' (Kho Tổng)
    # Dải size chuẩn từ 38 đến 47
    SIZES = [38, 39, 40, 41, 42, 43, 44, 45, 46, 47]

    sh_store = wb_val['Kho'] # Cửa hàng (Kho 2)
    sh_home = wb_val['Ở Nhà'] # Kho tổng (Kho 1)

    # Đọc tồn kho Cửa Hàng (Kho 2)
    # Header: cột 1 là Code, cột 2 là size 38, ..., cột 11 là size 47
    store_stock = {} # code -> {size: qty}
    for r in range(2, sh_store.max_row + 1):
        code = clean_code(sh_store.cell(row=r, column=1).value)
        if not code: continue
        store_stock[code] = {}
        for c_idx, sz in enumerate(SIZES, 2):
            val = sh_store.cell(row=r, column=c_idx).value
            qty = int(val) if isinstance(val, (int, float)) and val > 0 else 0
            store_stock[code][sz] = qty

    # Đọc tồn kho Kho Tổng Ở Nhà (Kho 1)
    home_stock = {} # code -> {size: qty}
    for r in range(2, sh_home.max_row + 1):
        code = clean_code(sh_home.cell(row=r, column=1).value)
        if not code: continue
        home_stock[code] = {}
        for c_idx, sz in enumerate(SIZES, 2):
            val = sh_home.cell(row=r, column=c_idx).value
            qty = int(val) if isinstance(val, (int, float)) and val > 0 else 0
            home_stock[code][sz] = qty

    # Danh sách tất cả mã giày duy nhất
    all_codes = sorted(list(set(list(store_stock.keys()) + list(home_stock.keys()))))
    print(f"-> Đã phát hiện tổng cộng {len(all_codes)} mã sản phẩm giày da Kingsman!")

    total_wh1_pairs = 0
    total_wh2_pairs = 0
    product_id_map = {} # code -> product_id
    variant_id_map = {} # (code, size) -> variant_id

    for code in all_codes:
        cat, name, color, mat = guess_category_and_info(code)
        cost = prices.get(code, {}).get('cost', 350000)
        retail = prices.get(code, {}).get('retail', 650000)

        cursor.execute("""
        INSERT INTO products (code, name, category, material, color, cost_price, retail_price, min_stock_per_size, notes)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, (code, name, cat, mat, color, cost, retail, 2, "Nhập từ file Excel Kingsman"))
        p_id = cursor.lastrowid
        product_id_map[code] = p_id

        # Tạo các biến thể size từ 38 đến 47
        for sz in SIZES:
            barcode = f"{code}-{sz}"
            cursor.execute("""
            INSERT INTO product_variants (product_id, size, barcode)
            VALUES (?, ?, ?);
            """, (p_id, sz, barcode))
            v_id = cursor.lastrowid
            variant_id_map[(code, sz)] = v_id

            # Tồn Kho 1 (Ở Nhà)
            q1 = home_stock.get(code, {}).get(sz, 0)
            # Tồn Kho 2 (Cửa Hàng)
            q2 = store_stock.get(code, {}).get(sz, 0)

            cursor.execute("INSERT INTO inventory (variant_id, warehouse_id, quantity) VALUES (?, 1, ?);", (v_id, q1))
            cursor.execute("INSERT INTO inventory (variant_id, warehouse_id, quantity) VALUES (?, 2, ?);", (v_id, q2))

            total_wh1_pairs += q1
            total_wh2_pairs += q2

    print(f"-> Đã nạp thành công tồn kho:")
    print(f"   + Kho 1 (Kho Tổng - Ở Nhà):   {total_wh1_pairs} đôi")
    print(f"   + Kho 2 (Cửa Hàng / Showroom): {total_wh2_pairs} đôi")
    print(f"   + Tổng tồn toàn hệ thống:      {total_wh1_pairs + total_wh2_pairs} đôi")

    # 4. Nạp lịch sử nhập hàng (sheet 'Nhập')
    if 'Nhập' in wb_val.sheetnames:
        sh_n = wb_val['Nhập']
        inbound_count = 0
        for r in range(2, sh_n.max_row + 1):
            dt_val = sh_n.cell(row=r, column=1).value
            code = clean_code(sh_n.cell(row=r, column=2).value)
            sz_val = sh_n.cell(row=r, column=3).value
            qty_val = sh_n.cell(row=r, column=4).value
            price_val = sh_n.cell(row=r, column=5).value
            wh_val = sh_n.cell(row=r, column=6).value

            if not code or not qty_val: continue
            try:
                sz = int(float(sz_val)) if sz_val else 40
                qty = int(float(qty_val))
                price = (float(price_val) * 1000 if float(price_val) < 10000 else float(price_val)) if price_val else 350000
            except:
                continue

            target_wh = 1 if str(wh_val).strip().lower() == 'ở nhà' else 2
            date_str = dt_val.strftime("%Y-%m-%d %H:%M:%S") if isinstance(dt_val, datetime) else "2026-05-04 00:00:00"
            tx_code = f"PN-{r:04d}"

            cursor.execute("""
            INSERT INTO transactions (code, type, target_warehouse_id, partner_name, total_quantity, total_amount, notes, created_at)
            VALUES (?, 'INBOUND', ?, 'Xưởng Sản Xuất Kingsman', ?, ?, ?, ?);
            """, (tx_code, target_wh, qty, qty * price, f"Nhập hàng mã {code} size {sz}", date_str))
            tx_id = cursor.lastrowid

            v_id = variant_id_map.get((code, sz))
            if v_id:
                cursor.execute("""
                INSERT INTO transaction_items (transaction_id, variant_id, size, quantity, unit_price)
                VALUES (?, ?, ?, ?, ?);
                """, (tx_id, v_id, sz, qty, price))
            inbound_count += 1
        print(f"-> Đã nạp thành công {inbound_count} dòng lịch sử Nhập hàng!")

    # 5. Nạp lịch sử bán hàng (sheet 'Xuất')
    if 'Xuất' in wb_val.sheetnames:
        sh_x = wb_val['Xuất']
        outbound_count = 0
        for r in range(2, sh_x.max_row + 1):
            dt_val = sh_x.cell(row=r, column=1).value
            code = clean_code(sh_x.cell(row=r, column=2).value)
            sz_val = sh_x.cell(row=r, column=3).value
            qty_val = sh_x.cell(row=r, column=4).value
            price_val = sh_x.cell(row=r, column=5).value
            wh_val = sh_x.cell(row=r, column=6).value
            cust_name = sh_x.cell(row=r, column=9).value or "Khách mua lẻ"
            phone = sh_x.cell(row=r, column=10).value or ""
            addr = sh_x.cell(row=r, column=11).value or ""
            platform = sh_x.cell(row=r, column=12).value or "Cửa hàng"

            if not code or not qty_val: continue
            try:
                sz = int(float(sz_val)) if sz_val else 40
                qty = int(float(qty_val))
                price = (float(price_val) * 1000 if float(price_val) < 10000 else float(price_val)) if price_val else 650000
            except:
                continue

            src_wh = 1 if str(wh_val).strip().lower() == 'ở nhà' else 2
            date_str = dt_val.strftime("%Y-%m-%d %H:%M:%S") if isinstance(dt_val, datetime) else "2026-05-05 00:00:00"
            tx_code = f"PX-{r:04d}"

            note_full = f"{platform} | {phone} | {addr}".strip(" |")

            cursor.execute("""
            INSERT INTO transactions (code, type, source_warehouse_id, partner_name, total_quantity, total_amount, notes, created_at)
            VALUES (?, 'OUTBOUND', ?, ?, ?, ?, ?, ?);
            """, (tx_code, src_wh, str(cust_name).strip(), qty, qty * price, note_full, date_str))
            tx_id = cursor.lastrowid

            v_id = variant_id_map.get((code, sz))
            if v_id:
                cursor.execute("""
                INSERT INTO transaction_items (transaction_id, variant_id, size, quantity, unit_price)
                VALUES (?, ?, ?, ?, ?);
                """, (tx_id, v_id, sz, qty, price))
            outbound_count += 1
        print(f"-> Đã nạp thành công {outbound_count} dòng lịch sử Bán hàng / Xuất kho!")

    conn.commit()
    cursor.execute("PRAGMA foreign_keys = ON;")
    conn.close()

    print("==================================================================")
    print("HOÀN TẤT NHẬP DỮ LIỆU TỪ EXCEL VÀO HỆ THỐNG THÀNH CÔNG RỰC RỠ!")
    print("==================================================================")

if __name__ == "__main__":
    run_import()
