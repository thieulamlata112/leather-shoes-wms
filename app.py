import os
import io
import socket
from datetime import datetime
from typing import List, Optional, Dict, Any
from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import HTMLResponse, FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import database

from contextlib import asynccontextmanager

@asynccontextmanager
async def lifespan(app: FastAPI):
    database.init_db()
    yield

app = FastAPI(title="Hệ Thống Quản Lý Kho Giày Da Trực Tuyến", version="1.0.0", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- Schemas ---
class NewProductRequest(BaseModel):
    code: str
    name: str
    category: str = "Oxford"
    material: str = "Da bò"
    color: str = "Đen"
    cost_price: float = 0
    retail_price: float = 0
    min_stock_per_size: int = 2
    notes: Optional[str] = ""
    sizes: List[int] = database.SUPPORTED_SIZES.copy()
    initial_stock_wh1: Dict[str, int] = {}
    initial_stock_wh2: Dict[str, int] = {}

class TransactionItemInput(BaseModel):
    variant_id: int
    size: int
    quantity: int
    unit_price: float = 0
    system_quantity: Optional[int] = None
    actual_quantity: Optional[int] = None

class CreateTransactionRequest(BaseModel):
    type: str # INBOUND, OUTBOUND, TRANSFER, AUDIT
    source_warehouse_id: Optional[int] = None
    target_warehouse_id: Optional[int] = None
    partner_name: Optional[str] = ""
    notes: Optional[str] = ""
    items: List[TransactionItemInput]

# --- Helper: Lấy địa chỉ IP mạng nội bộ ---
def get_local_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "127.0.0.1"

# --- API Endpoints ---

@app.get("/api/network-info")
def get_network_info():
    local_ip = get_local_ip()
    port = 8000
    return {
        "local_ip": local_ip,
        "port": port,
        "local_url": f"http://localhost:{port}",
        "network_url": f"http://{local_ip}:{port}",
        "message": f"Mở trình duyệt trên điện thoại (cùng mạng Wi-Fi) và truy cập: http://{local_ip}:{port}"
    }

@app.get("/api/warehouses")
def list_warehouses():
    return database.get_warehouses()

@app.get("/api/products")
def list_products(warehouse_id: Optional[int] = None):
    return database.get_all_products_with_stock(warehouse_id)

@app.get("/api/products/{product_id}")
def get_product(product_id: int):
    prod = database.get_product_by_id(product_id)
    if not prod:
        raise HTTPException(status_code=404, detail="Không tìm thấy mẫu giày này.")
    return prod

@app.post("/api/products")
def create_new_product(payload: NewProductRequest):
    try:
        new_prod = database.create_product(payload.dict())
        return {"success": True, "product": new_prod}
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Lỗi khi thêm mẫu giày: {str(e)}")

@app.get("/api/barcode/lookup")
def lookup_barcode(code: str):
    res = database.lookup_barcode(code)
    if not res:
        raise HTTPException(status_code=404, detail=f"Không tìm thấy thông tin cho mã: {code}")
    return res

@app.get("/api/dashboard")
def dashboard_stats():
    return database.get_dashboard_stats()

@app.post("/api/transactions")
def create_transaction(payload: CreateTransactionRequest):
    data = payload.dict()
    # Kiểm tra tính hợp lệ cơ bản
    t_type = data["type"].upper()
    if t_type == "INBOUND" and not data.get("target_warehouse_id"):
        raise HTTPException(status_code=400, detail="Vui lòng chọn Kho Nhận hàng.")
    if t_type == "OUTBOUND" and not data.get("source_warehouse_id"):
        raise HTTPException(status_code=400, detail="Vui lòng chọn Kho Xuất hàng.")
    if t_type == "TRANSFER":
        if not data.get("source_warehouse_id") or not data.get("target_warehouse_id"):
            raise HTTPException(status_code=400, detail="Vui lòng chọn cả Kho Gửi và Kho Nhận.")
        if data["source_warehouse_id"] == data["target_warehouse_id"]:
            raise HTTPException(status_code=400, detail="Kho Gửi và Kho Nhận không được trùng nhau.")
    if t_type == "AUDIT" and not data.get("target_warehouse_id"):
        raise HTTPException(status_code=400, detail="Vui lòng chọn Kho cần cân đối kiểm kê.")

    if not data.get("items") or len(data["items"]) == 0:
        raise HTTPException(status_code=400, detail="Danh sách mặt hàng trống.")

    try:
        res = database.execute_stock_transaction(data)
        return res
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi thực hiện giao dịch kho: {str(e)}")



class NLTransactionRequest(BaseModel):
    text: str
    action: Optional[str] = "create" # "parse" or "create"

@app.post("/api/natural-language-transaction")
def natural_language_transaction(req: NLTransactionRequest):
    """Xử lý giao dịch từ văn bản tự nhiên tiếng Việt."""
    try:
        from text_parser import parse_and_create_transaction, parse_natural_language
        if req.action == "parse":
            parsed = parse_natural_language(req.text)
            return {"success": True, "message": "Phân tích thành công", "data": parsed}
        else:
            result = parse_and_create_transaction(req.text)
            if result.get("success"):
                return {"success": True, "message": result["message"], "data": result.get("parsed", {})}
            else:
                return {"success": False, "error": result.get("error"), "parsed": result.get("parsed", {})}
    except Exception as e:
        return {"success": False, "error": str(e)}
@app.get("/api/transactions")
def list_transactions(limit: int = 50, type: Optional[str] = None):
    return database.get_transactions_list(limit, type)

@app.post("/api/transactions/{transaction_id}/cancel")
def cancel_transaction(transaction_id: int):
    """Hủy phiếu giao dịch và tự động hoàn trả tồn kho."""
    try:
        res = database.cancel_transaction(transaction_id)
        return res
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi khi hủy phiếu: {str(e)}")

@app.post("/api/import/re-sync")
def resync_from_excel():
    try:
        import import_kingsman
        import_kingsman.run_import()
        return {"success": True, "message": "Đã đồng bộ lại toàn bộ dữ liệu từ file Excel KingsMan (1).xlsx thành công!"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi khi đồng bộ: {str(e)}")

# --- Xuất Phiếu Xuất Hàng Excel (Sheet Xuất chuẩn KingsMan) ---
@app.get("/api/export/excel")
def export_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Xuất"

    # Header styling chuẩn sheet Xuất của KingsMan (1).xlsx: nền cam đào #F9CB9C, chữ đỏ đậm #FF0000
    header_fill = PatternFill(start_color="FFF9CB9C", end_color="FFF9CB9C", fill_type="solid")
    header_font = Font(name="Calibri", size=13, bold=True, color="FFFF0000")
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = [
        "Ngày Bán", "Mã Giày", "Size", "Số lượng", "Giá", "Kho",
        "Chuyển Khoản", "COD", "Tên Khách Hàng", "SDT", "Địa Chỉ",
        "Nền Tảng", "zalo", "FB", "Ghi chú", "Đã Giao", "Size Đúng", "Tất"
    ]

    ws.row_dimensions[1].height = 26
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin

    col_widths = {
        1: 13,   # Ngày Bán
        2: 12,   # Mã Giày
        3: 10,   # Size
        4: 10,   # Số lượng
        5: 12,   # Giá
        6: 12,   # Kho
        7: 15,   # Chuyển Khoản
        8: 15,   # COD
        9: 24,   # Tên Khách Hàng
        10: 16,  # SDT
        11: 45,  # Địa Chỉ
        12: 14,  # Nền Tảng
        13: 16,  # zalo
        14: 10,  # FB
        15: 18,  # Ghi chú
        16: 20,  # Đã Giao
        17: 14,  # Size Đúng
        18: 10   # Tất
    }
    for col_idx, width in col_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # Chỉ xuất các phiếu xuất hàng (OUTBOUND)
    tx_list = database.get_transactions_list(limit=10000, tx_type="OUTBOUND")
    current_row = 2

    for tx in tx_list:
        dt_str = tx.get("created_at") or ""
        try:
            sale_date = datetime.strptime(dt_str[:10], "%Y-%m-%d").date()
        except Exception:
            sale_date = datetime.now().date()

        # Kho: Kho 1 (Kho Tổng/Ở Nhà) -> 'Ở Nhà', Kho 2 (Cửa Hàng) -> 'Kho'
        src_wh = tx.get("source_warehouse_id")
        kho_str = "Ở Nhà" if src_wh == 1 else "Kho"

        # Tên khách hàng
        pname = tx.get("partner_name") or "Khách lẻ"
        if " - " in pname:
            pname = pname.split(" - ")[0].strip()

        # Thông tin SĐT, Địa chỉ, Nền tảng
        phone = tx.get("customer_phone") or ""
        address = tx.get("customer_address") or ""
        platform = tx.get("platform") or ""
        notes = tx.get("notes") or ""

        # Trích xuất fallback từ notes nếu chưa có
        if notes and "|" in notes:
            parts = [p.strip() for p in notes.split("|")]
            if any(":" in p for p in parts):
                for p in parts:
                    if ":" in p:
                        k, v = p.split(":", 1)
                        k = k.strip().lower()
                        v = v.strip()
                        if ("sđt" in k or "phone" in k or "sdt" in k) and not phone:
                            phone = v
                        elif ("đ/c" in k or "địa chỉ" in k or "dia chi" in k) and not address:
                            address = v
                        elif ("nền tảng" in k or "nen tang" in k) and (not platform or platform == "Khác"):
                            platform = v
            elif len(parts) >= 3:
                if not platform or platform == "Khác":
                    platform = parts[0]
                if not phone:
                    phone = parts[1]
                if not address:
                    address = parts[2]

        if not platform or platform == "Khác":
            platform = "Page"
        if not phone:
            phone = "K"
        if not address:
            address = "K"

        # Zalo & FB
        zalo_val = pname if ("zalo" in platform.lower() or "voz" in platform.lower()) else "K"
        fb_val = "K"

        # Ghi chú
        p_count = tx.get("purchase_count") or 1
        ghi_chu = f"Mua lần {p_count}" if p_count > 1 else "Mua lần 1"

        # Trạng thái giao hàng
        is_cancelled = tx.get("status") == "CANCELLED"
        da_giao = "ĐÃ HỦY ĐƠN" if is_cancelled else "Giao Thành Công"

        items = tx.get("items", [])
        if not items:
            tot_qty = tx.get("total_quantity", 1) or 1
            tot_amt = tx.get("total_amount", 0.0) or 0.0
            u_price = tot_amt / tot_qty if tot_qty else 0.0
            items = [{
                "product_code": "-",
                "size": "-",
                "quantity": tot_qty,
                "unit_price": u_price
            }]

        for it in items:
            ws.row_dimensions[current_row].height = 20
            p_code = it.get("product_code", "-")
            size_val = it.get("size", "-")
            qty_val = it.get("quantity", 1)

            # Quy đổi giá sang đơn vị nghìn đồng (890.000 -> 890.0) giống file KingsMan gốc
            raw_price = it.get("unit_price", 0.0) or 0.0
            if raw_price >= 10000:
                price_k = round(raw_price / 1000.0, 1)
            elif raw_price > 0:
                price_k = round(raw_price, 1)
            else:
                price_k = 0.0

            # Phân bổ Chuyển Khoản vs COD
            is_ck = any(k in notes.lower() for k in ["ck", "chuyển khoản", "chuyen khoan", "banking"])
            if is_ck:
                val_ck = round(price_k * qty_val, 1)
                val_cod = "K"
            else:
                val_ck = "K"
                val_cod = round(price_k * qty_val, 1)

            row_data = [
                (sale_date, "yyyy-mm-dd", "center"),
                (p_code, None, "center"),
                (size_val, None, "center"),
                (qty_val, None, "center"),
                (price_k, "#,##0.0", "center"),
                (kho_str, None, "center"),
                (val_ck, "#,##0.0" if isinstance(val_ck, (int, float)) else None, "center"),
                (val_cod, "#,##0.0" if isinstance(val_cod, (int, float)) else None, "center"),
                (pname, None, "center"),
                (phone, "@", "center"),
                (address, None, "left" if len(str(address)) > 25 else "center"),
                (platform, None, "center"),
                (zalo_val, None, "center"),
                (fb_val, None, "center"),
                (ghi_chu, None, "center"),
                (da_giao, None, "center"),
                ("", None, "center"),
                ("", None, "center")
            ]

            for col_idx, (val, num_fmt, align_h) in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(horizontal=align_h, vertical="center")
                cell.border = border_thin
                if num_fmt:
                    cell.number_format = num_fmt

            current_row += 1

    # Lưu vào buffer byte
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"KingsMan_Phieu_Xuat_Hang_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# --- Mount Static Frontend ---
STATIC_DIR = os.path.join(os.path.dirname(__file__), "static")
if not os.path.exists(STATIC_DIR):
    os.makedirs(STATIC_DIR, exist_ok=True)

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

@app.get("/")
def serve_index():
    index_path = os.path.join(STATIC_DIR, "index.html")
    if os.path.exists(index_path):
        return FileResponse(
            index_path,
            headers={
                "Cache-Control": "no-cache, no-store, must-revalidate",
                "Pragma": "no-cache",
                "Expires": "0"
            }
        )
    return HTMLResponse("<h1>Hệ Thống Quản Lý Kho Giày Da đang khởi động...</h1>")

if __name__ == "__main__":
    import uvicorn
    # Lắng nghe trên 0.0.0.0 để điện thoại cùng mạng Wi-Fi có thể truy cập được
    print("=" * 60)
    print("  HE THONG QUAN LY KHO GIAY DA DANG CHAY")
    print(f"  - May tinh:    http://localhost:8000")
    print(f"  - Dien thoai:  http://{get_local_ip()}:8000")
    print("=" * 60)
    uvicorn.run(app, host="0.0.0.0", port=8000)
